"""
End-to-End Engine & Report Generation Tests.
Creates a synthetic wiring diagram PDF with intentional non-compliances,
runs the QC Engine, and validates the output JSON, PDF, and XLSX reports.
"""

import io
import os
import tempfile
from reportlab.lib.pagesizes import letter
from reportlab.pdfgen import canvas
import openpyxl
from pypdf import PdfReader

from backend.src.ai.engine import QCAnalysisEngine
from backend.src.ai.schemas import OverallStatusEnum, SeverityEnum
from backend.src.reports.pdf_generator import PDFReportGenerator
from backend.src.reports.xlsx_generator import XLSXReportGenerator


def create_sample_schematic_pdf(filepath: str):
    """Generate a sample wiring diagram PDF with known discrepancies."""
    c = canvas.Canvas(filepath, pagesize=letter)

    # Title Block (Missing revision and date intentionally)
    c.setFont("Helvetica-Bold", 10)
    c.drawString(400, 50, "TITLE: INDUSTRIAL CONTROL HARNESS")
    c.drawString(400, 35, "DWG NO: WD-9002-A")
    c.drawString(400, 20, "DRAWN BY: Rajesh Kumar")
    # c.drawString(400, 5, "REV: B") <-- Intentionally missing!

    # Wiring details
    c.setFont("Helvetica", 9)
    # Wire 1: Valid
    c.drawString(72, 700, "W101 18 AWG RED (J1-P1)")
    # Wire 2: Missing Gauge (Violation!)
    c.drawString(72, 670, "W102 BLU (J1-P2)")
    # Wire 3: Missing Color (Violation!)
    c.drawString(72, 640, "W103 22 AWG (TB1-P3)")

    # Connectors
    c.drawString(72, 550, "CONNECTORS: J1 (MS3106A-20-29P), TB1, P1, P2, P3")

    # Notes
    c.drawString(72, 450, "GENERAL NOTES:")
    c.drawString(72, 435, "1. ALL WIRES TO BE TEFLON INSULATED.")
    c.drawString(72, 420, "2. STRIP LENGTH 5.0mm +/- 0.5mm.")

    c.save()


def test_qc_engine_full_run():
    with tempfile.TemporaryDirectory() as tmpdir:
        sample_pdf = os.path.join(tmpdir, "sample_harness.pdf")
        create_sample_schematic_pdf(sample_pdf)

        engine = QCAnalysisEngine()
        result = engine.analyze(sample_pdf, standards=["IPC-WHMA-A-620D", "UL 508A"])

        # Assertions
        assert result.filename == "sample_harness.pdf"
        assert result.overall_status == OverallStatusEnum.FAIL
        assert result.summary.checks_total >= 10
        assert result.summary.failed >= 2  # Missing gauge (critical) + missing color (major)

        # Check finding details
        finding_rule_ids = [f.rule_id for f in result.findings]
        assert "RULE-WG-001" in finding_rule_ids  # Missing gauge flagged
        assert "RULE-CC-003" in finding_rule_ids  # Missing color flagged
        assert "RULE-TB-005" in finding_rule_ids  # Incomplete title block flagged

        # Verify sequential finding IDs
        for idx, f in enumerate(result.findings, start=1):
            assert f.id == f"D-{idx:03d}"

        # Test PDF Report Generation
        pdf_out = os.path.join(tmpdir, "qc_report.pdf")
        pdf_gen = PDFReportGenerator()
        pdf_gen.generate(result, pdf_out)
        assert os.path.exists(pdf_out)
        assert os.path.getsize(pdf_out) > 1000

        # Verify PDF is readable
        reader = PdfReader(pdf_out)
        assert len(reader.pages) >= 1
        pdf_text = reader.pages[0].extract_text()
        assert "WIRING DIAGRAM QC ASSISTANT" in pdf_text
        assert "STATUS: FAIL" in pdf_text

        # Test XLSX Report Generation
        xlsx_out = os.path.join(tmpdir, "qc_report.xlsx")
        xlsx_gen = XLSXReportGenerator()
        xlsx_gen.generate(result, xlsx_out)
        assert os.path.exists(xlsx_out)
        assert os.path.getsize(xlsx_out) > 1000

        # Verify Excel sheets
        wb = openpyxl.load_workbook(xlsx_out)
        assert "QC Summary" in wb.sheetnames
        assert "Discrepancy Details" in wb.sheetnames
        summary_sheet = wb["QC Summary"]
        assert summary_sheet["A1"].value == "WIRING DIAGRAM QC ASSISTANT — AUDIT SUMMARY"
