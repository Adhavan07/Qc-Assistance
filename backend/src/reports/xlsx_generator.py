"""
Structured XLSX Report Generator.
Utilizes openpyxl to generate spreadsheet audits containing all findings and summary metrics.
"""

import io
from typing import Union
import openpyxl
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

from ..ai.schemas import QCAnalysisResult, SeverityEnum


class XLSXReportGenerator:
    """Compiles QC findings into formatted Excel workbooks for ERP/PLM integration."""

    HEADER_FILL = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    BOLD_FONT = Font(name="Calibri", size=11, bold=True)
    REGULAR_FONT = Font(name="Calibri", size=10)
    BORDER = Border(
        left=Side(style="thin", color="CBD5E1"),
        right=Side(style="thin", color="CBD5E1"),
        top=Side(style="thin", color="CBD5E1"),
        bottom=Side(style="thin", color="CBD5E1"),
    )

    def generate(self, result: QCAnalysisResult, output_destination: Union[str, io.BytesIO]):
        wb = openpyxl.Workbook()

        # Sheet 1: Summary
        ws_sum = wb.active
        ws_sum.title = "QC Summary"
        self._build_summary_sheet(ws_sum, result)

        # Sheet 2: Findings Details
        ws_findings = wb.create_sheet(title="Discrepancy Details")
        self._build_findings_sheet(ws_findings, result)

        wb.save(output_destination)

    def _build_summary_sheet(self, ws, result: QCAnalysisResult):
        ws.append(["WIRING DIAGRAM QC ASSISTANT — AUDIT SUMMARY"])
        ws["A1"].font = Font(name="Calibri", size=14, bold=True, color="0F172A")
        ws.append([])

        metadata_rows = [
            ("Document ID", result.document_id),
            ("Filename", result.filename),
            ("Overall Status", result.overall_status.value),
            ("Applied Standards", ", ".join(result.standards_applied)),
            ("Ruleset Version", result.rules_version),
            ("Model Version", result.model_version),
            ("Processing Duration (ms)", result.processing_time_ms),
        ]
        for label, val in metadata_rows:
            ws.append([label, val])
            row_idx = ws.max_row
            ws.cell(row=row_idx, column=1).font = self.BOLD_FONT
            ws.cell(row=row_idx, column=2).font = self.REGULAR_FONT

        ws.append([])
        ws.append(["Metric", "Count"])
        header_row = ws.max_row
        for col in range(1, 3):
            cell = ws.cell(row=header_row, column=col)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT

        summary_metrics = [
            ("Total Checks Performed", result.summary.checks_total),
            ("Passed Checks", result.summary.passed),
            ("Failed Checks", result.summary.failed),
            ("Review Required", result.summary.review),
            ("Critical Findings", result.summary.critical_count),
            ("Major Findings", result.summary.major_count),
            ("Minor Findings", result.summary.minor_count),
            ("Info Findings", result.summary.info_count),
        ]
        for metric, count in summary_metrics:
            ws.append([metric, count])
            r = ws.max_row
            ws.cell(row=r, column=1).border = self.BORDER
            ws.cell(row=r, column=2).border = self.BORDER

        self._autofit_columns(ws)

    def _build_findings_sheet(self, ws, result: QCAnalysisResult):
        headers = [
            "Finding ID",
            "Page",
            "Severity",
            "Category",
            "Description",
            "Evidence",
            "Requirement",
            "Standard",
            "Section",
            "Recommendation",
            "Confidence Level",
            "Confidence Score",
        ]
        ws.append(headers)
        for col_idx in range(1, len(headers) + 1):
            cell = ws.cell(row=1, column=col_idx)
            cell.fill = self.HEADER_FILL
            cell.font = self.HEADER_FONT
            cell.alignment = Alignment(horizontal="center")

        for f in result.findings:
            ws.append(
                [
                    f.id,
                    f.page,
                    f.severity.value,
                    f.category,
                    f.description,
                    f.evidence,
                    f.requirement,
                    f.standard,
                    f.standard_section,
                    f.recommendation,
                    f.confidence.level.value,
                    f.confidence.score,
                ]
            )
            r = ws.max_row
            for c in range(1, len(headers) + 1):
                cell = ws.cell(row=r, column=c)
                cell.font = self.REGULAR_FONT
                cell.border = self.BORDER
                if c == 3:  # Severity highlight
                    cell.font = Font(
                        name="Calibri",
                        size=10,
                        bold=True,
                        color="DC2626" if f.severity == SeverityEnum.CRITICAL else ("EA580C" if f.severity == SeverityEnum.MAJOR else "000000"),
                    )

        self._autofit_columns(ws)

    def _autofit_columns(self, ws):
        for col in ws.columns:
            max_len = max(len(str(cell.value or "")) for cell in col)
            col_letter = get_column_letter(col[0].column)
            ws.column_dimensions[col_letter].width = min(max(max_len + 3, 12), 45)
